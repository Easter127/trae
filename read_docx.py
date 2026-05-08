from docx import Document
import json

doc = Document('/workspace/小金量化模型分析研报_2026-05-07.docx')

print("=== Document Paragraphs ===\n")
for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)

print("\n\n=== Document Tables ===\n")
for i, table in enumerate(doc.tables):
    print(f"Table {i}:")
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        print(row_data)
